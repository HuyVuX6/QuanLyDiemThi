# -*- coding: utf-8 -*-
from odoo import models, fields, api, exceptions, Command
from datetime import timedelta
import re
import logging
import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

_logger = logging.getLogger(__name__)

class ExamExam(models.Model):
    _name = 'exam.exam'
    _description = 'Đề Thi'
    _rec_name = 'name'

    _sql_constraints = [
        ('code_unique', 'UNIQUE(code)', 'Mã đề thi phải duy nhất!'),
        ('total_score_pos', 'CHECK(total_score > 0)', 'Tổng điểm đề thi phải > 0!'),
    ]

    code = fields.Char(string='Mã Đề Thi', required=True, size=30)
    name = fields.Char(string='Tên Đề Thi', required=True, size=255)
    exam_type = fields.Selection([
        ('QuickTest', 'Kiểm Tra Thường Xuyên (15 Phút)'),
        ('MidTerm', 'Kiểm Tra Giữa Kỳ'),
        ('FinalTerm', 'Kiểm Tra Cuối Kỳ')
    ], string='Loại Đề Thi', default='QuickTest', required=True)
    topic = fields.Char(string='Chủ Đề Kiến Thức', default='Tổng hợp Toán lớp 3')
    status = fields.Selection([
        ('draft', 'Dự Thảo'),
        ('active', 'Đang Hoạt Động'),
        ('closed', 'Đã Đóng')
    ], string='Trạng Thái', default='active', required=True)
    auto_analyze = fields.Boolean(
        string='Tự Động Phân Tích AI',
        default=True,
        help='Tự động kích hoạt phân tích Claude AI ngay sau khi học sinh nộp bài'
    )
    total_score = fields.Float(string='Tổng Điểm', default=10.0, required=True)
    time_limit_minutes = fields.Integer(string='Thời Gian (Phút)', default=15)
    question_ids = fields.Many2many(
        comodel_name='exam.question',
        relation='exam_exam_question_rel',
        column1='exam_id',
        column2='question_id',
        string='Ngân Hàng Câu Hỏi'
    )
    total_questions = fields.Integer(
        string='Tổng Số Câu Hỏi',
        compute='_compute_total_questions',
        store=True
    )
    session_ids = fields.One2many('exam.session', 'exam_id', string='Các Lượt Làm Bài')

    @api.depends('question_ids')
    def _compute_total_questions(self):
        for record in self:
            record.total_questions = len(record.question_ids)

    def _compute_display_name(self):
        for record in self:
            record.display_name = f"[{record.code}] {record.name}"

    def action_open_assign_wizard(self):
        """Mở cửa sổ Wizard để Giáo viên gán đề thi cho cả lớp hoặc nhiều học sinh"""
        self.ensure_one()
        if self.status != 'active':
            raise exceptions.UserError("Đề thi này đã đóng hoặc chưa kích hoạt, không thể gán đề cho học sinh!")
        return {
            'name': f'Gán Đề Thi: {self.name}',
            'type': 'ir.actions.act_window',
            'res_model': 'exam.assign.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_exam_id': self.id,
            }
        }


class ExamQuestion(models.Model):
    _name = 'exam.question'
    _description = 'Câu Hỏi Toán Lớp 3'
    _rec_name = 'content'

    _sql_constraints = [
        ('code_unique', 'UNIQUE(code)', 'Mã câu hỏi phải duy nhất!'),
        ('points_pos', 'CHECK(points > 0)', 'Điểm câu hỏi phải > 0!'),
    ]

    code = fields.Char(string='Mã Câu Hỏi', required=True, size=30)
    content = fields.Text(string='Nội Dung Câu Hỏi', required=True)
    question_type = fields.Selection([
        ('numeric', 'Số Học / Đo Lường (Điền Số/Kết Quả)'),
        ('single_choice', 'Trắc Nghiệm (A, B, C, D)'),
        ('text', 'Tự Luận Ngắn / Khái Niệm')
    ], string='Dạng Câu Hỏi', default='numeric', required=True)
    correct_answer = fields.Char(string='Đáp Án Chuẩn', required=True)
    acceptable_answers = fields.Char(
        string='Đáp Án Chấp Nhận Thêm',
        help='Các đáp án tương đương được chấp nhận, phân tách bởi dấu | hoặc ; (ví dụ: 15|15 quả|15 quả táo)'
    )
    tolerance = fields.Float(
        string='Dung Sai Số Học',
        default=0.0,
        help='Dung sai cho phép khi so sánh kết quả số học'
    )
    points = fields.Float(string='Điểm', default=1.0, required=True)
    topic = fields.Selection([
        ('Addition', 'Phép Cộng'),
        ('Subtraction', 'Phép Trừ'),
        ('Multiplication', 'Phép Nhân'),
        ('Division', 'Phép Chia'),
        ('Geometry', 'Hình Học'),
        ('WordProblem', 'Toán Lời Văn'),
    ], string='Chủ Đề Toán', required=True)
    difficulty = fields.Selection([
        ('Easy', 'Dễ'),
        ('Medium', 'Trung Bình'),
        ('Hard', 'Khó')
    ], string='Độ Khó', default='Medium', required=True)
    explanation = fields.Text(string='Hướng Dẫn Giải / Lời Giải Chi Tiết')

    def _compute_display_name(self):
        for record in self:
            short_c = record.content[:40] if record.content else ''
            record.display_name = f"[{record.code}] {short_c}..."

    def action_open_import_wizard(self):
        """Mở Wizard Import Câu Hỏi Từ File Excel"""
        return {
            'name': 'Import Ngân Hàng Câu Hỏi Từ Excel',
            'type': 'ir.actions.act_window',
            'res_model': 'exam.question.import.wizard',
            'view_mode': 'form',
            'target': 'new',
        }


class ExamBatch(models.Model):
    _name = 'exam.batch'
    _description = 'Đợt Kiểm Tra Của Lớp'
    _order = 'date desc, id desc'
    _rec_name = 'name'

    _sql_constraints = [
        ('code_unique', 'UNIQUE(code)', 'Mã đợt kiểm tra phải là duy nhất!'),
    ]

    code = fields.Char(string='Mã Đợt', required=True, size=50)
    name = fields.Char(string='Tên Đợt Kiểm Tra', required=True, size=255)
    class_id = fields.Many2one('school.class', string='Lớp Học', required=True, ondelete='cascade', index=True)
    grade_id = fields.Many2one('school.grade', string='Khối Lớp', related='class_id.grade_id', store=True, readonly=True)
    batch_type = fields.Selection([
        ('entry', 'Kiểm Tra Đầu Kỳ / Khảo Sát'),
        ('midterm', 'Kiểm Tra Giữa Kỳ'),
        ('final', 'Kiểm Tra Cuối Kỳ'),
        ('regular', 'Kiểm Tra Thường Xuyên'),
        ('retest', 'Thi Lại / Bổ Sung (Dành Cho HS Rớt)'),
    ], string='Đợt Kiểm Tra', default='midterm', required=True)
    exam_id = fields.Many2one('exam.exam', string='Đề Thi Áp Dụng', required=True, domain=[('status', '=', 'active')], ondelete='restrict')
    date = fields.Date(string='Ngày Kiểm Tra', default=fields.Date.today, required=True)
    delivery_mode = fields.Selection([
        ('offline', 'Thi Offline (Làm bài trên giấy / Excel)'),
        ('online', 'Thi Trực Tuyến Trên Odoo'),
    ], string='Thể Thức Thi', default='offline', required=True)
    status = fields.Selection([
        ('draft', 'Dự Thảo'),
        ('in_progress', 'Đang Diễn Ra'),
        ('completed', 'Đã Hoàn Thành & Chấm Điểm'),
    ], string='Trạng Thái', default='draft', required=True)

    parent_batch_id = fields.Many2one('exam.batch', string='Thuộc Đợt Gốc', ondelete='set null', help='Đợt kiểm tra gốc nếu đây là đợt thi lại')
    retest_batch_ids = fields.One2many('exam.batch', 'parent_batch_id', string='Các Đợt Thi Lại')

    session_ids = fields.One2many('exam.session', 'batch_id', string='Danh Sách Bài Thi')

    total_students = fields.Integer(string='Tổng Học Sinh', compute='_compute_batch_stats', store=True)
    graded_count = fields.Integer(string='Đã Chấm', compute='_compute_batch_stats', store=True)
    passed_count = fields.Integer(string='Số Học Sinh Đạt', compute='_compute_batch_stats', store=True)
    failed_count = fields.Integer(string='Số Học Sinh Rớt', compute='_compute_batch_stats', store=True)
    pass_rate = fields.Float(string='Tỷ Lệ Đạt (%)', compute='_compute_batch_stats', store=True)
    average_score = fields.Float(string='Điểm Trung Bình', compute='_compute_batch_stats', store=True)
    note = fields.Text(string='Ghi Chú')

    @api.onchange('class_id', 'batch_type')
    def _onchange_suggest_name(self):
        if self.class_id and self.batch_type:
            type_dict = dict(self._fields['batch_type'].selection)
            type_name = type_dict.get(self.batch_type, 'Kiểm Tra')
            if not self.name:
                self.name = f"{type_name} - {self.class_id.name}"
            if not self.code:
                type_code_map = {'entry': 'DK', 'midterm': 'GK', 'final': 'CK', 'regular': 'TX', 'retest': 'TL'}
                prefix = type_code_map.get(self.batch_type, 'KT')
                self.code = f"{prefix}-{self.class_id.code or 'LOP'}-{fields.Date.today().strftime('%Y%m')}"

    @api.depends('session_ids', 'session_ids.status', 'session_ids.total_score', 'session_ids.passed')
    def _compute_batch_stats(self):
        for record in self:
            sessions = record.session_ids
            record.total_students = len(sessions)
            graded = sessions.filtered(lambda s: s.status == 'Graded')
            record.graded_count = len(graded)
            if graded:
                passed = graded.filtered(lambda s: s.passed)
                failed = graded.filtered(lambda s: not s.passed)
                record.passed_count = len(passed)
                record.failed_count = len(failed)
                record.pass_rate = round((len(passed) / len(graded)) * 100.0, 1)
                record.average_score = round(sum(s.total_score for s in graded) / len(graded), 2)
            else:
                record.passed_count = 0
                record.failed_count = 0
                record.pass_rate = 0.0
                record.average_score = 0.0

    def action_populate_students(self):
        """Tự động tạo lượt làm bài cho toàn bộ học sinh trong lớp"""
        self.ensure_one()
        if not self.class_id:
            raise exceptions.UserError("Vui lòng chọn lớp học trước khi lập danh sách học sinh!")
        if not self.exam_id:
            raise exceptions.UserError("Vui lòng chọn đề thi áp dụng trước khi lập danh sách học sinh!")
        if not self.class_id.student_ids:
            raise exceptions.UserError(f"Lớp {self.class_id.name} hiện chưa có học sinh nào. Vui lòng thêm học sinh vào lớp trước!")

        existing_student_ids = self.session_ids.mapped('student_id.id')
        new_sessions = []
        for student in self.class_id.student_ids:
            if student.id not in existing_student_ids:
                new_sessions.append({
                    'batch_id': self.id,
                    'student_id': student.id,
                    'class_id': self.class_id.id,
                    'exam_id': self.exam_id.id,
                    'status': 'NotStarted',
                })
        created_count = len(new_sessions)
        if new_sessions:
            self.env['exam.session'].create(new_sessions)
            if self.status == 'draft':
                self.status = 'in_progress'

        message = f"Đã lập danh sách {created_count} học sinh vào đợt kiểm tra!" if created_count else "Tất cả học sinh trong lớp đã có trong danh sách đợt thi."
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Thành công',
                'message': message,
                'type': 'success',
                'sticky': False,
            }
        }

    def action_create_retest_batch(self):
        """Tạo một đợt thi lại cho các học sinh bị rớt (điểm < 5.0) trong đợt này"""
        self.ensure_one()
        failed_sessions = self.session_ids.filtered(lambda s: s.status == 'Graded' and not s.passed)
        if not failed_sessions:
            raise exceptions.UserError(
                f"Đợt kiểm tra '{self.name}' hiện không có học sinh nào bị rớt (hoặc chưa chấm bài xong)!\n"
                f"Chỉ có thể tạo đợt thi lại cho học sinh có điểm dưới 5.0."
            )

        retest_count = len(self.retest_batch_ids) + 1
        retest_code = f"{self.code}-TL{retest_count}"
        retest_name = f"Thi Lại (Lần {retest_count}) - {self.name}"

        retest_batch = self.create({
            'code': retest_code,
            'name': retest_name,
            'class_id': self.class_id.id,
            'batch_type': 'retest',
            'exam_id': self.exam_id.id,
            'parent_batch_id': self.id,
            'date': fields.Date.today(),
            'delivery_mode': self.delivery_mode,
            'status': 'in_progress',
        })

        retest_sessions = []
        for s in failed_sessions:
            retest_sessions.append({
                'batch_id': retest_batch.id,
                'student_id': s.student_id.id,
                'class_id': self.class_id.id,
                'exam_id': self.exam_id.id,
                'status': 'NotStarted',
            })
        self.env['exam.session'].create(retest_sessions)

        return {
            'name': f"Đợt Thi Lại Cho Học Sinh Rớt: {retest_name}",
            'type': 'ir.actions.act_window',
            'res_model': 'exam.batch',
            'res_id': retest_batch.id,
            'view_mode': 'form',
            'target': 'current',
        }

    def action_open_import_wizard(self):
        """Mở wizard import bài làm thi offline từ file Excel"""
        self.ensure_one()
        return {
            'name': f"Import Bài Làm Excel - {self.name}",
            'type': 'ir.actions.act_window',
            'res_model': 'exam.batch.offline.import.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_batch_id': self.id,
            }
        }

    def action_download_offline_template(self):
        """Tải phiếu nhập bài làm Excel"""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': f"/educational_system/download_batch_offline_template/{self.id}",
            'target': 'self',
        }


class ExamSession(models.Model):
    _name = 'exam.session'
    _description = 'Lượt Làm Bài Của Học Sinh'
    _rec_name = 'name'

    _sql_constraints = [
        ('unique_student_batch', 'UNIQUE(student_id, batch_id)', 'Mỗi học sinh chỉ làm bài thi 1 lần trong cùng một đợt kiểm tra!'),
    ]

    name = fields.Char(compute='_compute_name', store=True, string='Tên Phiên')
    batch_id = fields.Many2one('exam.batch', string='Đợt Kiểm Tra', ondelete='cascade', index=True)
    student_id = fields.Many2one('school.student', string='Học Sinh', required=True, ondelete='cascade')
    student_code = fields.Char(related='student_id.code', string='Mã Sinh Viên / Mã HS', store=True, readonly=True)
    student_name = fields.Char(related='student_id.name', string='Họ Tên Học Sinh', store=True, readonly=True)
    exam_id = fields.Many2one(
        'exam.exam', 
        string='Đề Thi', 
        required=True, 
        domain=[('status', '=', 'active')],
        ondelete='cascade'
    )
    time_limit_minutes = fields.Integer(
        string='Thời Gian Làm Bài (Phút)', 
        related='exam_id.time_limit_minutes', 
        readonly=True
    )
    class_id = fields.Many2one(
        'school.class', 
        string='Lớp Học', 
        compute='_compute_class_id', 
        store=True, 
        readonly=False,
        ondelete='restrict'
    )
    start_time = fields.Datetime(string='Thời Gian Bắt Đầu')
    end_time = fields.Datetime(string='Thời Gian Nộp Bài')
    deadline = fields.Datetime(
        string='Hạn Chót Nộp Bài', 
        compute='_compute_deadline', 
        store=True,
        help='Thời điểm tối đa học sinh phải hoàn thành bài làm'
    )
    status = fields.Selection([
        ('NotStarted', 'Chưa Bắt Đầu'),
        ('InProgress', 'Đang Làm'),
        ('Submitted', 'Đã Nộp'),
        ('Graded', 'Đã Chấm'),
    ], default='NotStarted', string='Trạng Thái', required=True)
    auto_analyze = fields.Boolean(
        string='Tự Động Phân Tích AI',
        default=True,
        help='Tự động gửi sang mô hình AI phân tích ngay khi nộp bài'
    )

    total_possible = fields.Float(
        string='Tổng Điểm Tối Đa',
        related='exam_id.total_score',
        store=True
    )
    total_score = fields.Float(
        string='Điểm Đạt Được',
        compute='_compute_scores',
        store=True
    )
    correct_count = fields.Integer(
        string='Số Câu Đúng',
        compute='_compute_scores',
        store=True
    )
    wrong_count = fields.Integer(
        string='Số Câu Sai',
        compute='_compute_scores',
        store=True
    )
    percentage = fields.Float(
        string='Tỷ Lệ (%)',
        compute='_compute_scores',
        store=True
    )
    passed = fields.Boolean(
        string='Kết Quả Đạt',
        compute='_compute_passed',
        store=True
    )
    result_display = fields.Char(
        string='Kết Quả',
        compute='_compute_passed',
        store=True
    )

    answer_ids = fields.One2many('exam.answer', 'session_id', string='Chi Tiết Bài Làm')
    analysis_ids = fields.One2many('ai.analysis', 'session_id', string='Kết Quả AI')

    @api.constrains('exam_id')
    def _check_exam_status(self):
        for record in self:
            if record.exam_id and record.exam_id.status != 'active':
                raise exceptions.ValidationError(
                    f"Đề thi '{record.exam_id.name}' đã đóng hoặc chưa được kích hoạt! Vui lòng chỉ chọn đề thi đang ở trạng thái 'Đang Hoạt Động'."
                )

    @api.depends('start_time', 'time_limit_minutes')
    def _compute_deadline(self):
        for record in self:
            if record.start_time and record.time_limit_minutes:
                record.deadline = record.start_time + timedelta(minutes=record.time_limit_minutes)
            else:
                record.deadline = False

    @api.depends('student_id.class_id')
    def _compute_class_id(self):
        for record in self:
            if record.student_id and record.student_id.class_id:
                record.class_id = record.student_id.class_id

    @api.onchange('student_id')
    def _onchange_student_id(self):
        if self.student_id and self.student_id.class_id:
            self.class_id = self.student_id.class_id

    @api.onchange('exam_id')
    def _onchange_exam_id(self):
        """Khi chọn đề thi, tự động điền danh sách câu hỏi của đề thi vào bài làm"""
        self.answer_ids = [Command.clear()]
        if self.exam_id:
            lines = [
                Command.create({
                    'question_id': question.id,
                    'answer_given': '',
                })
                for question in self.exam_id.question_ids
            ]
            self.answer_ids = lines

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            # 1. Tự động lấy class_id từ học sinh nếu chưa có
            if vals.get('student_id') and not vals.get('class_id'):
                student = self.env['school.student'].browse(vals['student_id'])
                if student.class_id:
                    vals['class_id'] = student.class_id.id

            # 2. Tự động nạp bộ câu hỏi từ exam_id nếu chưa có answer_ids
            if vals.get('exam_id') and not vals.get('answer_ids'):
                exam = self.env['exam.exam'].browse(vals['exam_id'])
                vals['answer_ids'] = [
                    Command.create({'question_id': q.id, 'answer_given': ''})
                    for q in exam.question_ids
                ]
        return super().create(vals_list)

    @api.constrains('start_time', 'end_time')
    def _check_exam_times(self):
        for record in self:
            if record.start_time and record.end_time and record.end_time < record.start_time:
                raise exceptions.ValidationError("Thời gian nộp bài phải sau thời gian bắt đầu làm bài!")

    @api.depends('student_id.name', 'exam_id.name')
    def _compute_name(self):
        for record in self:
            record.name = f"{record.student_id.name or ''} - {record.exam_id.name or ''}"

    @api.depends('status', 'answer_ids.points_earned', 'answer_ids.is_correct', 'total_possible', 'answer_ids')
    def _compute_scores(self):
        for record in self:
            answers = record.answer_ids
            total_q = len(answers)
            # Học sinh không được nhìn thấy Tỷ lệ, Số câu đúng/sai, Điểm đạt được khi đang làm bài
            if record.status != 'Graded':
                record.total_score = 0.0
                record.correct_count = 0
                record.wrong_count = 0
                record.percentage = 0.0
                continue

            correct_answers = answers.filtered(lambda a: a.is_correct)
            record.correct_count = len(correct_answers)
            record.wrong_count = total_q - record.correct_count

            if total_q > 0 and record.total_possible and record.total_possible > 0:
                # Phương thức tính điểm chuẩn: Điểm mỗi câu = total_possible / total_questions (Ví dụ: 10 / 40 = 0.25)
                point_per_q = record.total_possible / total_q
                record.total_score = round(record.correct_count * point_per_q, 2)
                record.percentage = round((record.correct_count / total_q) * 100.0, 1)
            else:
                record.total_score = round(sum(ans.points_earned for ans in answers), 2)
                if record.total_possible and record.total_possible > 0:
                    record.percentage = round((record.total_score / record.total_possible) * 100.0, 1)
                else:
                    record.percentage = 0.0

    def action_start_exam(self):
        """Bắt đầu tính giờ làm bài cho học sinh"""
        for record in self:
            if record.status != 'NotStarted':
                raise exceptions.UserError("Bài thi này đã bắt đầu hoặc đã hoàn thành!")
            record.write({
                'start_time': fields.Datetime.now(),
                'status': 'InProgress',
            })

    def action_submit_exam(self):
        self.ensure_one()
        if not self.start_time:
            self.start_time = fields.Datetime.now()
        self.end_time = fields.Datetime.now()
        self.status = 'Graded'
        
        # Đảm bảo chấm lại điểm toàn bộ bài thi
        for answer in self.answer_ids:
            answer._compute_evaluation()
        self._compute_scores()

        # Tạo hoặc lấy bản ghi AI Analysis
        analysis = self.env['ai.analysis'].search([('session_id', '=', self.id)], limit=1)
        if not analysis:
            analysis = self.env['ai.analysis'].create({
                'session_id': self.id,
                'student_id': self.student_id.id,
                'status': 'pending',
            })

        # Tự động gọi AI nếu được kích hoạt
        if self.auto_analyze:
            try:
                analysis.action_analyze_with_claude()
            except Exception as e:
                _logger.warning(f"Không thể tự động phân tích AI cho phiên {self.id}: {str(e)}")

    @api.depends('total_score', 'status')
    def _compute_passed(self):
        for record in self:
            if record.status == 'Graded':
                record.passed = (record.total_score >= 5.0)
                record.result_display = 'Đạt' if record.passed else 'Chưa Đạt (Rớt)'
            else:
                record.passed = False
                record.result_display = 'Chưa Chấm'

    def action_view_session_detail(self):
        """Mở chi tiết bài thi của học sinh để xem kết quả và phân tích AI"""
        self.ensure_one()
        return {
            'name': f"Chi Tiết Bài Làm - {self.student_name} ({self.exam_id.name})",
            'type': 'ir.actions.act_window',
            'res_model': 'exam.session',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'current',
        }

    @api.model
    def cron_auto_submit_expired_sessions(self):
        """Tự động nộp bài và chấm điểm các phiên thi đã hết hạn làm bài"""
        now = fields.Datetime.now()
        expired_sessions = self.search([
            ('status', '=', 'InProgress'),
            ('deadline', '<=', now),
        ])
        for session in expired_sessions:
            _logger.info(f"Cron: Tự động nộp bài cho phiên thi quá hạn {session.id} ({session.name})")
            try:
                session.action_submit_exam()
            except Exception as e:
                _logger.error(f"Lỗi tự động nộp bài cho phiên {session.id}: {str(e)}")


class ExamAnswer(models.Model):
    _name = 'exam.answer'
    _description = 'Chi Tiết Đáp Án'

    _sql_constraints = [
        ('unique_session_question', 'UNIQUE(session_id, question_id)', 'Một câu hỏi chỉ được trả lời 1 lần trong một lượt làm bài!'),
    ]

    session_id = fields.Many2one('exam.session', string='Phiên Thi', required=True, ondelete='cascade')
    question_id = fields.Many2one('exam.question', string='Câu Hỏi', required=True, ondelete='cascade')
    
    # Các trường liên kết để học sinh tiện đọc đề bài trên cùng một dòng
    question_content = fields.Text(string='Đề Bài', related='question_id.content', readonly=True)
    question_type = fields.Selection(string='Dạng Câu Hỏi', related='question_id.question_type', readonly=True)
    question_points = fields.Float(
        string='Thang Điểm', 
        compute='_compute_question_points', 
        store=True,
        help='Thang điểm chuẩn của câu hỏi (tổng điểm / số câu)'
    )
    
    answer_given = fields.Char(string='Đáp Án Của Học Sinh')
    is_correct = fields.Boolean(string='Đúng/Sai', compute='_compute_evaluation', store=True)
    points_earned = fields.Float(string='Điểm Nhận Được', compute='_compute_evaluation', store=True)

    @api.depends('session_id.total_possible', 'session_id.answer_ids')
    def _compute_question_points(self):
        for record in self:
            sess = record.session_id
            total_possible = sess.total_possible if sess and sess.total_possible else 10.0
            total_q = len(sess.answer_ids) if sess else 0
            if total_q > 0:
                record.question_points = round(total_possible / total_q, 2)
            elif record.question_id and record.question_id.points:
                record.question_points = record.question_id.points
            else:
                record.question_points = 1.0

    @api.constrains('question_id', 'session_id')
    def _check_question_belongs_to_exam(self):
        """Đảm bảo câu hỏi bắt buộc phải nằm trong Đề thi của lượt thi này!"""
        for record in self:
            if record.session_id and record.question_id and record.session_id.exam_id:
                if record.question_id not in record.session_id.exam_id.question_ids:
                    raise exceptions.ValidationError(
                        f"Lỗi: Câu hỏi '{record.question_id.code}' không thuộc Đề thi '{record.session_id.exam_id.name}'!"
                    )

    @staticmethod
    def _normalize_choice(val):
        if not val:
            return ""
        m = re.search(r'\b([A-Da-d])\b', str(val).strip())
        if m:
            return m.group(1).upper()
        return str(val).strip().upper()

    @staticmethod
    def _extract_number(val):
        if not val:
            return None
        val_clean = str(val).strip().replace(',', '.')
        m = re.search(r'[-+]?\d+(?:\.\d+)?', val_clean)
        if m:
            try:
                return float(m.group(0))
            except ValueError:
                return None
        return None

    @api.depends('answer_given', 'question_id.correct_answer', 'question_points',
                 'question_id.question_type', 'question_id.acceptable_answers', 'question_id.tolerance',
                 'session_id.status')
    def _compute_evaluation(self):
        for record in self:
            # Chỉ hiển thị kết quả đúng/sai và điểm khi bài thi đã được chấm (Graded)
            if not record.session_id or record.session_id.status != 'Graded':
                record.is_correct = False
                record.points_earned = 0.0
                continue

            q = record.question_id
            if not record.answer_given or not q or not q.correct_answer:
                record.is_correct = False
                record.points_earned = 0.0
                continue

            s_raw = str(record.answer_given).strip()
            c_raw = str(q.correct_answer).strip()
            pts = record.question_points if record.question_points > 0 else (q.points or 1.0)

            # 1. So khớp trực tiếp (bỏ qua hoa thường)
            if s_raw.lower() == c_raw.lower():
                record.is_correct = True
                record.points_earned = pts
                continue

            # 2. So khớp với danh sách đáp án chấp nhận (ngăn cách bởi | hoặc ;)
            all_correct = [c_raw]
            for source in [q.correct_answer, q.acceptable_answers]:
                if source:
                    for alt in re.split(r'[|;]', str(source)):
                        alt_clean = alt.strip()
                        if alt_clean and alt_clean not in all_correct:
                            all_correct.append(alt_clean)

            if any(s_raw.lower() == alt.lower() for alt in all_correct):
                record.is_correct = True
                record.points_earned = pts
                continue

            # 3. So khớp theo dạng câu hỏi
            is_matched = False
            q_type = q.question_type or 'numeric'

            if q_type == 'single_choice':
                s_choice = self._normalize_choice(s_raw)
                c_choice = self._normalize_choice(c_raw)
                if s_choice and c_choice and s_choice == c_choice:
                    is_matched = True

            elif q_type == 'numeric':
                s_num = self._extract_number(s_raw)
                c_num = self._extract_number(c_raw)
                if s_num is not None and c_num is not None:
                    tol = q.tolerance or 0.0
                    if abs(s_num - c_num) <= tol:
                        is_matched = True

            if is_matched:
                record.is_correct = True
                record.points_earned = pts
            else:
                record.is_correct = False
                record.points_earned = 0.0


class ExamAssignWizard(models.TransientModel):
    _name = 'exam.assign.wizard'
    _description = 'Wizard Gán Đề Thi Cho Học Sinh'

    exam_id = fields.Many2one('exam.exam', string='Đề Thi Cần Gán', required=True, readonly=True)
    assign_type = fields.Selection([
        ('class', 'Gán Cho Toàn Bộ Lớp Học'),
        ('students', 'Chọn Từng Học Sinh'),
    ], string='Hình Thức Gán', default='class', required=True)
    class_id = fields.Many2one('school.class', string='Lớp Học')
    student_ids = fields.Many2many('school.student', string='Danh Sách Học Sinh')

    @api.constrains('exam_id')
    def _check_exam_active(self):
        for wizard in self:
            if wizard.exam_id and wizard.exam_id.status != 'active':
                raise exceptions.ValidationError("Đề thi này đã đóng hoặc chưa kích hoạt, không thể phát đề!")

    @api.onchange('class_id')
    def _onchange_class_id(self):
        if self.class_id:
            return {'domain': {'student_ids': [('class_id', '=', self.class_id.id)]}}

    def action_assign_exam(self):
        """Thực hiện phát đề thi: Tạo sẵn các lượt làm bài kèm câu hỏi cho học sinh"""
        self.ensure_one()
        students = self.env['school.student']
        if self.assign_type == 'class':
            if not self.class_id:
                raise exceptions.UserError("Vui lòng chọn lớp học để gán đề thi!")
            students = self.class_id.student_ids.filtered(lambda s: s.status == 'Active')
            if not students:
                raise exceptions.UserError(f"Lớp {self.class_id.name} chưa có học sinh nào đang học!")
        else:
            if not self.student_ids:
                raise exceptions.UserError("Vui lòng chọn ít nhất 1 học sinh!")
            students = self.student_ids

        created_sessions = self.env['exam.session']
        skipped_names = []

        for student in students:
            # Kiểm tra xem học sinh đã có lượt làm bài cho đề này chưa
            existing = self.env['exam.session'].search([
                ('student_id', '=', student.id),
                ('exam_id', '=', self.exam_id.id)
            ], limit=1)
            if existing:
                skipped_names.append(student.name)
                continue

            # Tạo lượt làm bài (câu hỏi sẽ tự động được sinh ra trong create())
            session = self.env['exam.session'].create({
                'student_id': student.id,
                'exam_id': self.exam_id.id,
                'class_id': student.class_id.id if student.class_id else False,
                'status': 'NotStarted',
            })
            created_sessions |= session

        msg = f"Đã gán đề thi '{self.exam_id.name}' cho {len(created_sessions)} học sinh thành công!"
        if skipped_names:
            msg += f"\n(Bỏ qua {len(skipped_names)} học sinh đã được gán trước đó: {', '.join(skipped_names)})"

        action = self.env['ir.actions.act_window']._for_xml_id('educational_system.action_exam_session')
        action.update({
            'name': f'Lượt Làm Bài - {self.exam_id.name}',
            'domain': [('exam_id', '=', self.exam_id.id)],
            'context': {'default_exam_id': self.exam_id.id},
        })
        return action


class ExamQuestionImportWizard(models.TransientModel):
    _name = 'exam.question.import.wizard'
    _description = 'Wizard Import Ngân Hàng Câu Hỏi Từ Excel'

    excel_file = fields.Binary(string='Tệp Excel (.xlsx)', required=False)
    filename = fields.Char(string='Tên Tệp')
    duplicate_behavior = fields.Selection([
        ('update', 'Cập nhật nội dung câu hỏi nếu mã đã tồn tại (Khuyên dùng)'),
        ('skip', 'Bỏ qua dòng trùng mã (giữ nguyên câu hỏi cũ)'),
        ('error', 'Báo lỗi và dừng import khi phát hiện trùng mã'),
    ], string='Quy Tắc Xử Lý Trùng Mã', default='update', required=True)
    default_points = fields.Float(string='Điểm Mặc Định Nếu Trống', default=1.0, required=True)
    state = fields.Selection([
        ('draft', 'Chọn Tệp'),
        ('done', 'Hoàn Tất'),
    ], default='draft')
    result_html = fields.Html(string='Kết QuẢ Import', readonly=True)

    def action_download_template(self):
        """Trả về URL tải về file mẫu trực tiếp qua controller, không lo bị chặn bởi validate form"""
        return {
            'type': 'ir.actions.act_url',
            'url': '/educational_system/download_question_template',
            'target': 'self',
        }

    def action_import_excel(self):
        """Đọc tệp Excel và import câu hỏi theo bộ nguyên tắc chuẩn hóa & kiểm tra lỗi"""
        self.ensure_one()
        if not self.excel_file:
            raise exceptions.UserError("Vui lòng chọn tệp Excel (.xlsx) để tải lên!")

        if self.filename and not (self.filename.lower().endswith('.xlsx') or self.filename.lower().endswith('.xls')):
            raise exceptions.UserError("Định dạng tệp không hợp lệ! Vui lòng chỉ chọn tệp Excel có đuôi .xlsx hoặc .xls.")

        try:
            file_data = base64.b64decode(self.excel_file)
            wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
        except Exception as e:
            raise exceptions.UserError(f"Không thể đọc file Excel. Vui lòng đảm bảo file không bị khóa hoặc lỗi định dạng! Chi tiết: {str(e)}")

        sheet = wb.active
        if not sheet or sheet.max_row < 2:
            raise exceptions.UserError("File Excel rỗng hoặc không chứa dòng dữ liệu nào sau dòng tiêu đề!")

        header_aliases = {
            'code': ['mã câu hỏi', 'mã ch', 'mã', 'code', 'ma cau hoi', 'mach'],
            'content': ['nội dung câu hỏi', 'nội dung', 'đề bài', 'câu hỏi', 'content', 'noi dung'],
            'topic': ['chủ đề', 'dạng toán', 'dạng bài', 'chủ đề toán', 'topic', 'chu de', 'dang toan'],
            'question_type': ['dạng câu hỏi', 'loại câu hỏi', 'kiểu câu hỏi', 'question_type', 'dang cau hoi', 'loai cau hoi'],
            'difficulty': ['độ khó', 'mức độ', 'difficulty', 'do kho', 'muc do'],
            'points': ['điểm', 'thang điểm', 'points', 'diem'],
            'correct_answer': ['đáp án chuẩn', 'đáp án đúng', 'đáp án', 'correct_answer', 'dap an', 'dap an chuan'],
            'acceptable_answers': ['đáp án chấp nhận thêm', 'đáp án tương đương', 'đáp án phụ', 'acceptable_answers', 'dap an phu'],
            'tolerance': ['dung sai', 'dung sai số học', 'tolerance'],
            'explanation': ['hướng dẫn giải', 'lời giải', 'giải thích', 'lời giải chi tiết', 'explanation', 'huong dan giai', 'loi giai'],
        }

        col_mapping = {}
        header_row = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]

        for col_idx, raw_val in enumerate(header_row, start=1):
            if not raw_val:
                continue
            clean_hdr = str(raw_val).strip().lower().replace('(*)', '').strip()
            for field_name, aliases in header_aliases.items():
                if any(alias in clean_hdr for alias in aliases):
                    if field_name not in col_mapping:
                        col_mapping[field_name] = col_idx
                    break

        required_cols = ['code', 'content', 'topic', 'correct_answer']
        missing_required = [col for col in required_cols if col not in col_mapping]
        if missing_required:
            col_names_vi = {
                'code': 'Mã Câu Hỏi',
                'content': 'Nội Dung Câu Hỏi',
                'topic': 'Chủ Đề',
                'correct_answer': 'Đáp Án Chuẩn',
            }
            missing_vi = [col_names_vi.get(c, c) for c in missing_required]
            raise exceptions.UserError(
                f"File Excel thiếu các cột bắt buộc: {', '.join(missing_vi)}.\n"
                f"Vui lòng tải file mẫu để kiểm tra đúng cấu trúc các cột!"
            )

        total_rows = 0
        created_count = 0
        updated_count = 0
        skipped_count = 0
        errors = []
        seen_codes_in_file = set()

        for row_idx in range(2, sheet.max_row + 1):
            row_data = {}
            has_any_value = False
            for field_name, col_idx in col_mapping.items():
                val = sheet.cell(row=row_idx, column=col_idx).value
                if val is not None and str(val).strip() != '':
                    has_any_value = True
                    row_data[field_name] = str(val).strip()
                else:
                    row_data[field_name] = ''

            if not has_any_value:
                continue

            total_rows += 1

            code = row_data.get('code', '').strip().upper()
            if not code:
                errors.append(f"Dòng {row_idx}: Thiếu Mã Câu Hỏi.")
                continue

            if code in seen_codes_in_file:
                errors.append(f"Dòng {row_idx}: Trùng mã '{code}' với một dòng khác ngay trong tệp Excel.")
                continue
            seen_codes_in_file.add(code)

            content = row_data.get('content', '').strip()
            if not content:
                errors.append(f"Dòng {row_idx}: Thiếu Nội Dung Câu Hỏi (Mã: {code}).")
                continue

            correct_answer = row_data.get('correct_answer', '').strip()
            if not correct_answer:
                errors.append(f"Dòng {row_idx}: Thiếu Đáp Án Chuẩn (Mã: {code}).")
                continue

            topic_raw = row_data.get('topic', '').strip().lower()
            topic = None
            if any(k in topic_raw for k in ['cộng', 'addition', 'cong']):
                topic = 'Addition'
            elif any(k in topic_raw for k in ['trừ', 'subtraction', 'tru']):
                topic = 'Subtraction'
            elif any(k in topic_raw for k in ['nhân', 'multiplication', 'nhan']):
                topic = 'Multiplication'
            elif any(k in topic_raw for k in ['chia', 'division']):
                topic = 'Division'
            elif any(k in topic_raw for k in ['hình', 'geometry', 'hinh']):
                topic = 'Geometry'
            elif any(k in topic_raw for k in ['lời văn', 'văn', 'word', 'loi van']):
                topic = 'WordProblem'

            if not topic:
                errors.append(
                    f"Dòng {row_idx}: Chủ đề '{row_data.get('topic')}' không hợp lệ. "
                    f"Hỗ trợ: Phép Cộng, Phép Trừ, Phép Nhân, Phép Chia, Hình Học, Toán Lời Văn."
                )
                continue

            qtype_raw = row_data.get('question_type', '').strip().lower()
            if any(k in qtype_raw for k in ['trắc nghiệm', 'choice', 'trac nghiem']):
                question_type = 'single_choice'
            elif any(k in qtype_raw for k in ['tự luận', 'text', 'tu luan']):
                question_type = 'text'
            else:
                question_type = 'numeric'

            diff_raw = row_data.get('difficulty', '').strip().lower()
            if any(k in diff_raw for k in ['dễ', 'easy', 'de']):
                difficulty = 'Easy'
            elif any(k in diff_raw for k in ['khó', 'hard', 'kho']):
                difficulty = 'Hard'
            else:
                difficulty = 'Medium'

            try:
                points_val = float(row_data.get('points')) if row_data.get('points') else self.default_points
                points = points_val if points_val > 0 else 1.0
            except Exception:
                points = self.default_points if self.default_points > 0 else 1.0

            try:
                tolerance_val = float(row_data.get('tolerance')) if row_data.get('tolerance') else 0.0
                tolerance = tolerance_val if tolerance_val >= 0 else 0.0
            except Exception:
                tolerance = 0.0

            acceptable_answers = row_data.get('acceptable_answers', '').strip()
            explanation = row_data.get('explanation', '').strip()

            vals = {
                'code': code,
                'content': content,
                'topic': topic,
                'question_type': question_type,
                'difficulty': difficulty,
                'points': points,
                'correct_answer': correct_answer,
                'acceptable_answers': acceptable_answers or False,
                'tolerance': tolerance,
                'explanation': explanation or False,
            }

            existing = self.env['exam.question'].search([('code', '=', code)], limit=1)
            if existing:
                if self.duplicate_behavior == 'error':
                    raise exceptions.UserError(f"Phát hiện trùng mã câu hỏi '{code}' tại dòng {row_idx} với câu hỏi đã có trong hệ thống!")
                elif self.duplicate_behavior == 'skip':
                    skipped_count += 1
                    continue
                elif self.duplicate_behavior == 'update':
                    existing.write(vals)
                    updated_count += 1
            else:
                self.env['exam.question'].create(vals)
                created_count += 1

        err_html = ""
        if errors:
            err_items = "".join([f"<li class='text-danger mb-1'>{err}</li>" for err in errors])
            err_html = f"""
            <div class="alert alert-danger mt-3" role="alert">
                <h6 class="alert-heading fw-bold mb-2"><i class="fa fa-exclamation-triangle me-1"/> Các Dòng Bị Lỗi ({len(errors)} dòng không được import):</h6>
                <ul class="mb-0 ps-3">
                    {err_items}
                </ul>
            </div>
            """

        result_html = f"""
        <div class="p-2">
            <div class="alert alert-success" role="status">
                <h5 class="alert-heading fw-bold mb-3"><i class="fa fa-check-circle me-1"/> Kết Quả Import File Excel:</h5>
                <div class="row text-center g-2 mb-2">
                    <div class="col-3">
                        <div class="p-2 border rounded bg-white">
                            <small class="text-muted d-block">Tổng số dòng</small>
                            <span class="fs-4 fw-bold text-dark">{total_rows}</span>
                        </div>
                    </div>
                    <div class="col-3">
                        <div class="p-2 border rounded bg-white">
                            <small class="text-muted d-block">Tạo mới</small>
                            <span class="fs-4 fw-bold text-success">{created_count}</span>
                        </div>
                    </div>
                    <div class="col-3">
                        <div class="p-2 border rounded bg-white">
                            <small class="text-muted d-block">Cập nhật</small>
                            <span class="fs-4 fw-bold text-primary">{updated_count}</span>
                        </div>
                    </div>
                    <div class="col-3">
                        <div class="p-2 border rounded bg-white">
                            <small class="text-muted d-block">Bỏ qua / Lỗi</small>
                            <span class="fs-4 fw-bold text-warning">{skipped_count + len(errors)}</span>
                        </div>
                    </div>
                </div>
            </div>
            {err_html}
        </div>
        """

        self.write({
            'state': 'done',
            'result_html': result_html,
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_reset(self):
        """Cho phép tải lên tệp khác"""
        self.write({
            'state': 'draft',
            'excel_file': False,
            'filename': False,
            'result_html': False,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_view_questions(self):
        """Chuyển đến danh sách Ngân hàng câu hỏi"""
        return self.env['ir.actions.act_window']._for_xml_id('educational_system.action_exam_question')


class ExamBatchOfflineImportWizard(models.TransientModel):
    _name = 'exam.batch.offline.import.wizard'
    _description = 'Wizard Import Bài Làm Thi Off Từ Excel'

    batch_id = fields.Many2one('exam.batch', string='Đợt Kiểm Tra', required=True)
    excel_file = fields.Binary(string='Tệp Excel Bài Làm Học Sinh (*)', required=True)
    filename = fields.Char(string='Tên Tệp Excel')
    auto_analyze = fields.Boolean(
        string='Tự Động Kích Hoạt Phân Tích AI',
        default=True,
        help='Tự động phân tích điểm mạnh, điểm yếu qua Claude AI sau khi chấm'
    )
    state = fields.Selection([
        ('draft', 'Chọn Tệp'),
        ('done', 'Kết Quả Chấm Điểm')
    ], default='draft')
    result_html = fields.Html(string='Báo Cáo Kết Quả')

    def action_import_offline_results(self):
        self.ensure_one()
        if not self.excel_file:
            raise exceptions.UserError("Vui lòng chọn tệp Excel bài làm để tải lên!")

        batch = self.batch_id
        if not batch.exam_id:
            raise exceptions.UserError("Đợt kiểm tra chưa gán đề thi! Không thể chấm điểm.")

        exam = batch.exam_id
        questions = exam.question_ids
        if not questions:
            raise exceptions.UserError(f"Đề thi '{exam.name}' không có câu hỏi nào để chấm điểm!")

        try:
            file_data = base64.b64decode(self.excel_file)
            wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
            ws = wb.active
        except Exception as e:
            raise exceptions.UserError(f"Không thể đọc tệp Excel. Vui lòng đảm bảo tệp định dạng .xlsx hợp lệ! Lỗi: {str(e)}")

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise exceptions.UserError("Tệp Excel không có dữ liệu bài làm (cần ít nhất 1 dòng tiêu đề và 1 dòng học sinh)!")

        headers = [str(h).strip() if h is not None else '' for h in rows[0]]

        student_code_col = None
        student_name_col = None
        question_col_map = {}

        question_code_map = {q.code.strip().upper(): q for q in questions if q.code}

        for idx, h_text in enumerate(headers):
            h_upper = h_text.upper()
            if any(term in h_upper for term in ['MÃ HỌC SINH', 'MA HOC SINH', 'MÃ SINH VIÊN', 'MA SINH VIEN', 'STUDENT_CODE', 'MÃ HS']):
                student_code_col = idx
            elif any(term in h_upper for term in ['HỌ VÀ TÊN', 'HO VA TEN', 'HỌ TÊN', 'TÊN HỌC SINH', 'STUDENT_NAME']):
                student_name_col = idx
            else:
                match = re.search(r'\[\s*([A-Za-z0-9_-]+)\s*\]', h_text)
                if match:
                    code_extracted = match.group(1).strip().upper()
                    if code_extracted in question_code_map:
                        question_col_map[idx] = question_code_map[code_extracted]
                else:
                    clean_h = h_text.strip().upper()
                    if clean_h in question_code_map:
                        question_col_map[idx] = question_code_map[clean_h]

        if student_code_col is None:
            student_code_col = 0

        if not question_col_map:
            raise exceptions.UserError(
                f"Không tìm thấy cột câu hỏi nào khớp với đề thi '{exam.name}'!\n"
                f"Vui lòng sử dụng file mẫu tải về từ đợt kiểm tra (chứa mã câu hỏi trong tiêu đề, vd: [CH-T3-001])."
            )

        graded_students = []
        errors = []

        for row_idx, row in enumerate(rows[1:], start=2):
            raw_code = row[student_code_col] if student_code_col < len(row) else None
            if raw_code is None or str(raw_code).strip() == '':
                continue
            student_code = str(raw_code).strip()

            student = self.env['school.student'].search([('code', '=', student_code)], limit=1)
            if not student and batch.class_id:
                student = self.env['school.student'].search([
                    ('class_id', '=', batch.class_id.id),
                    ('code', '=ilike', student_code)
                ], limit=1)

            if not student and student_name_col is not None and student_name_col < len(row):
                raw_name = row[student_name_col]
                if raw_name:
                    student = self.env['school.student'].search([
                        ('class_id', '=', batch.class_id.id),
                        ('name', '=ilike', str(raw_name).strip())
                    ], limit=1)

            if not student:
                errors.append(f"Dòng {row_idx}: Không tìm thấy học sinh có mã '{student_code}'.")
                continue

            session = self.env['exam.session'].search([
                ('batch_id', '=', batch.id),
                ('student_id', '=', student.id)
            ], limit=1)

            if not session:
                session = self.env['exam.session'].create({
                    'batch_id': batch.id,
                    'student_id': student.id,
                    'class_id': batch.class_id.id,
                    'exam_id': exam.id,
                    'status': 'NotStarted',
                    'auto_analyze': self.auto_analyze,
                })
            else:
                session.auto_analyze = self.auto_analyze

            for col_idx, question in question_col_map.items():
                ans_val = ''
                if col_idx < len(row) and row[col_idx] is not None:
                    ans_val = str(row[col_idx]).strip()

                answer_line = session.answer_ids.filtered(lambda a: a.question_id.id == question.id)
                if answer_line:
                    answer_line.write({'answer_given': ans_val})
                else:
                    self.env['exam.answer'].create({
                        'session_id': session.id,
                        'question_id': question.id,
                        'answer_given': ans_val,
                    })

            session.action_submit_exam()
            graded_students.append({
                'code': student.code or '',
                'name': student.name or '',
                'score': session.total_score,
                'passed': session.passed,
                'correct': session.correct_count,
                'wrong': session.wrong_count,
            })

        batch.write({'status': 'completed'})

        total_graded = len(graded_students)
        passed_students = [s for s in graded_students if s['passed']]
        failed_students = [s for s in graded_students if not s['passed']]
        avg_score = round(sum(s['score'] for s in graded_students) / total_graded, 2) if total_graded > 0 else 0.0

        err_html = ""
        if errors:
            err_items = "".join(f"<li>{e}</li>" for e in errors)
            err_html = f"""
            <div class="alert alert-warning mt-3">
                <h6 class="fw-bold mb-1"><i class="fa fa-exclamation-triangle me-1"/> Cảnh báo / Dòng bỏ qua ({len(errors)}):</h6>
                <ul class="mb-0 ps-3 small">{err_items}</ul>
            </div>
            """

        student_rows_html = ""
        for s in graded_students:
            badge_class = "bg-success" if s['passed'] else "bg-danger"
            badge_text = "Đạt" if s['passed'] else "Chưa Đạt (Rớt)"
            score_class = "text-success" if s['passed'] else "text-danger"
            student_rows_html += f"""
            <tr>
                <td class="text-center fw-bold">{s['code']}</td>
                <td>{s['name']}</td>
                <td class="text-center fw-bold fs-6 {score_class}">{s['score']}</td>
                <td class="text-center text-success">{s['correct']}</td>
                <td class="text-center text-danger">{s['wrong']}</td>
                <td class="text-center"><span class="badge {badge_class}">{badge_text}</span></td>
            </tr>
            """

        result_html = f"""
        <div class="p-2">
            <div class="alert alert-success d-flex align-items-center justify-content-between p-3 mb-3">
                <div>
                    <h5 class="fw-bold mb-1"><i class="fa fa-check-circle me-1"/> Đã Chấm Điểm Xong Cho Đợt Kiểm Tra: {batch.name}</h5>
                    <p class="mb-0 small text-muted">Thể thức thi offline qua Excel. Toàn bộ bài thi đã được chấm điểm tự động và gửi phân tích AI.</p>
                </div>
            </div>

            <div class="row text-center g-2 mb-3">
                <div class="col-3">
                    <div class="p-2 border rounded bg-light">
                        <small class="text-muted d-block">Học Sinh Đã Chấm</small>
                        <span class="fs-4 fw-bold text-dark">{total_graded}</span>
                    </div>
                </div>
                <div class="col-3">
                    <div class="p-2 border rounded bg-light">
                        <small class="text-muted d-block">Điểm Trung Bình</small>
                        <span class="fs-4 fw-bold text-primary">{avg_score}</span>
                    </div>
                </div>
                <div class="col-3">
                    <div class="p-2 border rounded bg-light">
                        <small class="text-muted d-block">Số Học Sinh Đạt (&gt;= 5.0)</small>
                        <span class="fs-4 fw-bold text-success">{len(passed_students)}</span>
                    </div>
                </div>
                <div class="col-3">
                    <div class="p-2 border rounded bg-light">
                        <small class="text-muted d-block">Số Học Sinh Rớt (&lt; 5.0)</small>
                        <span class="fs-4 fw-bold text-danger">{len(failed_students)}</span>
                    </div>
                </div>
            </div>

            {err_html}

            <div class="table-responsive border rounded bg-white mt-3" style="max-height: 350px; overflow-y: auto;">
                <table class="table table-sm table-striped table-hover mb-0">
                    <thead class="table-primary sticky-top">
                        <tr>
                            <th class="text-center">Mã HS</th>
                            <th>Họ Và Tên</th>
                            <th class="text-center">Điểm Đạt Được</th>
                            <th class="text-center">Số Đúng</th>
                            <th class="text-center">Số Sai</th>
                            <th class="text-center">Kết Quả</th>
                        </tr>
                    </thead>
                    <tbody>
                        {student_rows_html}
                    </tbody>
                </table>
            </div>
        </div>
        """

        self.write({
            'state': 'done',
            'result_html': result_html,
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_reset(self):
        """Cho phép tải lên tệp khác"""
        self.write({
            'state': 'draft',
            'excel_file': False,
            'filename': False,
            'result_html': False,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_view_batch(self):
        """Quay lại xem chi tiết đợt kiểm tra"""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'exam.batch',
            'res_id': self.batch_id.id,
            'view_mode': 'form',
            'target': 'current',
        }