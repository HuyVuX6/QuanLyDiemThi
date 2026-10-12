# -*- coding: utf-8 -*-
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from odoo import http
from odoo.http import request


class ExamQuestionTemplateController(http.Controller):

    @http.route('/educational_system/download_question_template', type='http', auth='user')
    def download_question_template(self, **kwargs):
        """Tạo và trả về trực tiếp file Excel mẫu cho Ngân hàng câu hỏi Toán 3"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'CauHoiToan3'

        navy_header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC'),
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

        headers = [
            'Mã Câu Hỏi (*)',
            'Nội Dung Câu Hỏi (*)',
            'Chủ Đề (*)',
            'Dạng Câu Hỏi',
            'Độ Khó',
            'Điểm',
            'Đáp Án Chuẩn (*)',
            'Đáp Án Chấp Nhận Thêm',
            'Dung Sai',
            'Hướng Dẫn Giải',
        ]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        sample_rows = [
            ['CH-T3-001', 'Tính: 345 + 128 = ?', 'Phép Cộng', 'Điền số', 'Dễ', 1.0, '473', '', 0.0, '345 + 128 = 473'],
            ['CH-T3-002', 'Tính: 890 - 456 = ?', 'Phép Trừ', 'Điền số', 'Dễ', 1.0, '434', '', 0.0, '890 - 456 = 434'],
            ['CH-T3-003', 'Tính: 48 : 6 = ?', 'Phép Chia', 'Điền số', 'Dễ', 1.0, '8', '', 0.0, '48 : 6 = 8'],
            ['CH-T3-004', 'Một hình chữ nhật có chiều dài 8 cm, chiều rộng 5 cm. Tính chu vi của hình chữ nhật đó.', 'Hình Học', 'Điền số', 'Trung Bình', 1.0, '26', '26 cm|26cm', 0.0, 'Chu vi hình chữ nhật là: (8 + 5) x 2 = 26 (cm).'],
            ['CH-T3-005', 'Một cửa hàng có 5 bao gạo, mỗi bao nặng 9 kg. Hỏi cửa hàng có tất cả bao nhiêu ki-lô-gam gạo?', 'Toán Lời Văn', 'Điền số', 'Trung Bình', 1.0, '45', '45 kg|45kg', 0.0, 'Số gạo cửa hàng có là: 5 x 9 = 45 (kg).'],
            ['CH-T3-006', 'Số liền sau của số 999 là số nào?\nA. 998\nB. 1000\nC. 1001\nD. 990', 'Phép Cộng', 'Trắc nghiệm', 'Dễ', 1.0, 'B', '1000', 0.0, 'Số liền sau của 999 là 1000. Đáp án đúng là B.'],
        ]

        for row_data in sample_rows:
            ws.append(row_data)
            current_row = ws.max_row
            ws.row_dimensions[current_row].height = 24
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = Font(name='Arial', size=10)
                cell.border = thin_border
                if col_idx in [1, 3, 4, 5, 6, 7, 9]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        column_widths = {'A': 16, 'B': 45, 'C': 16, 'D': 15, 'E': 14, 'F': 10, 'G': 18, 'H': 25, 'I': 12, 'J': 35}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        ws_guide = wb.create_sheet(title='HuongDan_NguyenTac')
        ws_guide.column_dimensions['A'].width = 25
        ws_guide.column_dimensions['B'].width = 65

        guide_header = ['Quy Tắc / Mục', 'Chi Tiết Hướng Dẫn & Giá Trị Hợp Lệ']
        ws_guide.append(guide_header)
        for col_idx in range(1, 3):
            cell = ws_guide.cell(row=1, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align

        guide_rows = [
            ['Cột bắt buộc (*)', 'Mã Câu Hỏi, Nội Dung Câu Hỏi, Chủ Đề, Đáp Án Chuẩn.'],
            ['Chủ Đề hợp lệ', 'Phép Cộng (Addition), Phép Trừ (Subtraction), Phép Nhân (Multiplication), Phép Chia (Division), Hình Học (Geometry), Toán Lời Văn (WordProblem).'],
            ['Dạng Câu Hỏi', 'Điền số (numeric - mặc định), Trắc nghiệm (single_choice), Tự luận (text).'],
            ['Độ Khó', 'Dễ (Easy), Trung Bình (Medium - mặc định), Khó (Hard).'],
            ['Đáp Án Chuẩn', 'Đáp án chính xác tuyệt đối (ví dụ: 45 hoặc B).'],
            ['Đáp Án Phụ', 'Các cách viết tương đương, phân tách bởi dấu gạch đứng | hoặc dấu chấm phẩy ; (ví dụ: 45 kg|45kg).'],
            ['Xử Lý Trùng Mã', 'Tùy chọn: Cập nhật nếu đã có (khuyên dùng) hoặc Bỏ qua dòng trùng.'],
        ]
        for g_row in guide_rows:
            ws_guide.append(g_row)
            r = ws_guide.max_row
            ws_guide.cell(row=r, column=1).font = Font(name='Arial', size=10, bold=True)
            ws_guide.cell(row=r, column=2).font = Font(name='Arial', size=10)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        file_content = buf.getvalue()

        return request.make_response(
            file_content,
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', 'attachment; filename="mau_import_ngan_hang_cau_hoi_toan_3.xlsx"'),
            ]
        )

    @http.route('/educational_system/download_batch_offline_template/<int:batch_id>', type='http', auth='user')
    def download_batch_offline_template(self, batch_id, **kwargs):
        """Tạo và tải về file Excel mẫu nhập bài làm thi offline cho đợt kiểm tra"""
        batch = request.env['exam.batch'].browse(batch_id)
        if not batch.exists():
            return request.not_found()

        exam = batch.exam_id
        questions = exam.question_ids.sorted(key=lambda q: q.code or str(q.id))

        students = []
        if batch.session_ids:
            students = batch.session_ids.mapped('student_id')
        elif batch.class_id:
            students = batch.class_id.student_ids

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'BaiLamHocSinh'

        navy_header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC'),
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

        headers = ['Mã Học Sinh (*)', 'Họ Và Tên Học Sinh']
        for q in questions:
            clean_content = (q.content or '').replace('\n', ' ')
            if len(clean_content) > 35:
                clean_content = clean_content[:32] + '...'
            headers.append(f"[{q.code}] {clean_content}")

        ws.append(headers)
        ws.row_dimensions[1].height = 32

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        for s in students:
            row_data = [s.code or '', s.name or '']
            row_data.extend([''] * len(questions))
            ws.append(row_data)
            curr_row = ws.max_row
            ws.row_dimensions[curr_row].height = 24
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = Font(name='Arial', size=10)
                cell.border = thin_border
                if col_idx in [1]:
                    cell.alignment = center_align
                elif col_idx in [2]:
                    cell.alignment = left_align
                else:
                    cell.alignment = center_align

        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 28
        for col_idx in range(3, len(headers) + 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 25

        ws_info = wb.create_sheet(title='ThongTinDeThi')
        info_headers = ['Mã Câu Hỏi', 'Nội Dung Đầy Đủ', 'Dạng Câu Hỏi', 'Chủ Đề', 'Đáp Án Chuẩn', 'Thang Điểm']
        ws_info.append(info_headers)
        for col_idx in range(1, len(info_headers) + 1):
            c = ws_info.cell(row=1, column=col_idx)
            c.fill = navy_header_fill
            c.font = header_font
            c.alignment = center_align

        for q in questions:
            ws_info.append([
                q.code or '',
                q.content or '',
                q.question_type or '',
                q.topic or '',
                q.correct_answer or '',
                q.points or 1.0,
            ])
            r = ws_info.max_row
            for col_idx in range(1, len(info_headers) + 1):
                cell = ws_info.cell(row=r, column=col_idx)
                cell.font = Font(name='Arial', size=10)
                cell.border = thin_border
                if col_idx in [1, 3, 4, 5, 6]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        ws_info.column_dimensions['A'].width = 16
        ws_info.column_dimensions['B'].width = 50
        ws_info.column_dimensions['C'].width = 16
        ws_info.column_dimensions['D'].width = 16
        ws_info.column_dimensions['E'].width = 18
        ws_info.column_dimensions['F'].width = 12

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_content = output.read()

        filename = f"Phieu_Nhap_Bai_Thi_Off_{batch.code}.xlsx"
        return request.make_response(
            file_content,
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename="{filename}"')
            ]
        )

